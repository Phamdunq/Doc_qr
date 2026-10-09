import cv2
import os
import shutil
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# ============================================================
# CHƯƠNG TRÌNH ĐỌC QR CCCD TỪ NHIỀU ẢNH
#
# Cấu trúc thư mục:
#   chuong_trinh_doc_qr/
#       doc_qr_cccd.py
#       anh_cccd/                 <- bỏ ảnh vào đây
#       anh_khong_doc_duoc_qr/    <- ảnh không đọc được sẽ được copy vào đây
#
# Kết quả:
#   ket_qua_qr_cccd.xlsx
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "anh_cccd"
FAILED_DIR = BASE_DIR / "anh_khong_doc_duoc_qr"
OUTPUT_XLSX = BASE_DIR / "ket_qua_qr_cccd.xlsx"

IMAGE_EXTENSIONS = {
    ".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tif", ".tiff"
}

def parse_cccd_qr(data: str):
    """
    QR CCCD thường có dạng:
    CCCD||Họ tên|Ngày sinh|Giới tính|Địa chỉ|Ngày cấp

    Trả về 6 trường chính.
    Nếu định dạng khác, vẫn giữ toàn bộ dữ liệu ở cột QR_Goc.
    """
    data = data.strip().replace("\x00", "")

    parts = data.split("|")

    # Một số QR CCCD có || sau số CCCD.
    if len(parts) >= 7 and parts[1] == "":
        cccd = parts[0]
        ho_ten = parts[2]
        ngay_sinh = parts[3]
        gioi_tinh = parts[4]
        dia_chi = parts[5]
        ngay_cap = parts[6]
        return cccd, ho_ten, ngay_sinh, gioi_tinh, dia_chi, ngay_cap

    # Trường hợp dữ liệu dùng đúng 6 phần.
    if len(parts) >= 6:
        return (
            parts[0], parts[1], parts[2],
            parts[3], parts[4], parts[5]
        )

    return "", "", "", "", "", ""


def decode_once(detector, image):
    """Thử đọc QR trực tiếp từ một ảnh."""
    try:
        data, points, _ = detector.detectAndDecode(image)
        if data:
            return data
    except Exception:
        pass
    return None


def decode_qr_robust(image):
    """
    Đọc QR qua nhiều phiên bản ảnh để tăng khả năng đọc:
    - ảnh gốc
    - ảnh xám
    - tăng tương phản
    - threshold OTSU
    - threshold adaptive
    - làm nét
    - phóng to / thu nhỏ
    - xoay 90/180/270 độ

    Không phụ thuộc vào kích thước ảnh.
    """
    detector = cv2.QRCodeDetector()

    h, w = image.shape[:2]

    # Giới hạn kích thước để tránh xử lý ảnh quá lớn.
    # Ảnh nhỏ sẽ được phóng lên; ảnh lớn sẽ được thu xuống.
    scales = [1.0]

    if min(h, w) < 1000:
        scales.extend([1.5, 2.0, 3.0])
    elif max(h, w) > 2500:
        scales.extend([0.75, 0.5])

    candidates = []

    for scale in scales:
        if scale == 1.0:
            img = image
        else:
            img = cv2.resize(
                image, None, fx=scale, fy=scale,
                interpolation=cv2.INTER_CUBIC if scale > 1 else cv2.INTER_AREA
            )

        candidates.append(img)

        gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        candidates.append(gray)

        # Tăng tương phản cục bộ.
        clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
        enhanced = clahe.apply(gray)
        candidates.append(enhanced)

        # OTSU.
        _, otsu = cv2.threshold(
            enhanced, 0, 255,
            cv2.THRESH_BINARY + cv2.THRESH_OTSU
        )
        candidates.append(otsu)

        # Adaptive threshold.
        adaptive = cv2.adaptiveThreshold(
            enhanced, 255,
            cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
            cv2.THRESH_BINARY,
            31, 7
        )
        candidates.append(adaptive)

        # Làm nét nhẹ.
        blur = cv2.GaussianBlur(gray, (0, 0), 1.2)
        sharp = cv2.addWeighted(gray, 1.7, blur, -0.7, 0)
        candidates.append(sharp)

    # Thử ảnh gốc trước.
    for candidate in candidates:
        data = decode_once(detector, candidate)
        if data:
            return data

    # Thử xoay ảnh trong trường hợp ảnh bị xoay.
    for angle in (90, 180, 270):
        if angle == 90:
            rotated = cv2.rotate(image, cv2.ROTATE_90_CLOCKWISE)
        elif angle == 180:
            rotated = cv2.rotate(image, cv2.ROTATE_180)
        else:
            rotated = cv2.rotate(image, cv2.ROTATE_90_COUNTERCLOCKWISE)

        data = decode_once(detector, rotated)
        if data:
            return data

        gray = cv2.cvtColor(rotated, cv2.COLOR_BGR2GRAY)
        data = decode_once(detector, gray)
        if data:
            return data

    return None


def create_excel(rows, failed_files):
    wb = Workbook()

    ws = wb.active
    ws.title = "Du lieu QR"

    headers = [
        "STT",
        "Ten anh",
        "So CCCD",
        "Ho va ten",
        "Ngay sinh",
        "Gioi tinh",
        "Dia chi",
        "Ngay cap",
        "QR_Goc",
        "Trang thai"
    ]
    ws.append(headers)

    for i, row in enumerate(rows, start=1):
        ws.append([
            i,
            row["Ten anh"],
            row["So CCCD"],
            row["Ho va ten"],
            row["Ngay sinh"],
            row["Gioi tinh"],
            row["Dia chi"],
            row["Ngay cap"],
            row["QR_Goc"],
            row["Trang thai"]
        ])

    # Sheet ảnh lỗi.
    ws2 = wb.create_sheet("Anh khong doc duoc")
    ws2.append(["STT", "Ten anh", "Ly do", "Duong dan anh da copy"])

    for i, item in enumerate(failed_files, start=1):
        ws2.append([
            i,
            item["Ten anh"],
            item["Ly do"],
            item["Duong dan"]
        ])

    # Định dạng.
    for sheet in (ws, ws2):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        sheet.freeze_panes = "A2"

        for column_cells in sheet.columns:
            max_len = 0
            column_letter = column_cells[0].column_letter
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            sheet.column_dimensions[column_letter].width = min(max(max_len + 2, 12), 45)

    # Cột địa chỉ / QR gốc cho rộng hơn.
    ws.column_dimensions["G"].width = 45
    ws.column_dimensions["I"].width = 70
    ws2.column_dimensions["D"].width = 45

    wb.save(OUTPUT_XLSX)


def main():
    INPUT_DIR.mkdir(exist_ok=True)
    FAILED_DIR.mkdir(exist_ok=True)

    image_files = sorted(
        p for p in INPUT_DIR.iterdir()
        if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS
    )

    if not image_files:
        print("=" * 60)
        print("KHONG TIM THAY ANH")
        print(f"Hay bo anh vao thu muc: {INPUT_DIR}")
        print("=" * 60)
        input("Nhan Enter de thoat...")
        return

    rows = []
    failed_files = []

    print("=" * 60)
    print("        CHUONG TRINH DOC QR CCCD")
    print("=" * 60)
    print(f"So anh can doc: {len(image_files)}")
    print()

    for index, image_path in enumerate(image_files, start=1):
        print(f"[{index}/{len(image_files)}] Dang doc: {image_path.name}")

        image = cv2.imread(str(image_path))

        if image is None:
            print("   -> KHONG MO DUOC ANH")
            destination = FAILED_DIR / image_path.name
            shutil.copy2(image_path, destination)
            failed_files.append({
                "Ten anh": image_path.name,
                "Ly do": "Khong mo duoc anh",
                "Duong dan": str(destination)
            })
            continue

        data = decode_qr_robust(image)

        if data:
            cccd, ho_ten, ngay_sinh, gioi_tinh, dia_chi, ngay_cap = parse_cccd_qr(data)

            rows.append({
                "Ten anh": image_path.name,
                "So CCCD": cccd,
                "Ho va ten": ho_ten,
                "Ngay sinh": ngay_sinh,
                "Gioi tinh": gioi_tinh,
                "Dia chi": dia_chi,
                "Ngay cap": ngay_cap,
                "QR_Goc": data,
                "Trang thai": "Doc thanh cong"
            })

            print(f"   -> OK: {cccd} | {ho_ten}")
        else:
            print("   -> KHONG TIM THAY / KHONG DOC DUOC QR")

            destination = FAILED_DIR / image_path.name
            shutil.copy2(image_path, destination)

            failed_files.append({
                "Ten anh": image_path.name,
                "Ly do": "Khong co QR hoac khong giai ma duoc QR",
                "Duong dan": str(destination)
            })

    create_excel(rows, failed_files)

    print()
    print("=" * 60)
    print("HOAN TAT")
    print(f"Doc thanh cong : {len(rows)} anh")
    print(f"Anh loi        : {len(failed_files)} anh")
    print(f"File Excel     : {OUTPUT_XLSX}")
    print(f"Anh loi        : {FAILED_DIR}")
    print("=" * 60)
    input("Nhan Enter de thoat...")


if __name__ == "__main__":
    main()
